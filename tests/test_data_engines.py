"""Tests der Daten-Engines - sie laufen ohne externe Programme."""

import json

import pytest

from app.engines.base import ConversionError
from app.engines.data import TableEngine, TreeEngine

BEISPIEL_CSV = (
    "Aktenzeichen,Antragsteller,Eingang,Betrag\n"
    "AZ-2024-001,Müller GmbH,2024-03-01,1250.50\n"
    "AZ-2024-002,Schmidt KG,2024-03-04,890.00\n"
)


@pytest.fixture
def tabelle():
    return TableEngine()


@pytest.fixture
def baum():
    return TreeEngine()


def test_csv_nach_json(tabelle, tmp_path, ctx):
    quelle = tmp_path / "faelle.csv"
    quelle.write_text(BEISPIEL_CSV, encoding="utf-8")
    ziel = tmp_path / "faelle.json"

    tabelle.convert(quelle, ziel, "csv", "json", ctx)

    daten = json.loads(ziel.read_text(encoding="utf-8"))
    assert len(daten) == 2
    assert daten[0]["Antragsteller"] == "Müller GmbH"
    assert daten[1]["Aktenzeichen"] == "AZ-2024-002"


def test_csv_nach_xlsx_und_zurueck(tabelle, tmp_path, ctx):
    quelle = tmp_path / "faelle.csv"
    quelle.write_text(BEISPIEL_CSV, encoding="utf-8")
    mappe = tmp_path / "faelle.xlsx"
    zurueck = tmp_path / "zurueck.csv"

    tabelle.convert(quelle, mappe, "csv", "xlsx", ctx)
    assert mappe.stat().st_size > 0
    tabelle.convert(mappe, zurueck, "xlsx", "csv", ctx)

    inhalt = zurueck.read_text(encoding="utf-8-sig")
    assert "Müller GmbH" in inhalt
    assert "AZ-2024-002" in inhalt


def test_semikolon_wird_erkannt(tabelle, tmp_path, ctx):
    quelle = tmp_path / "deutsch.csv"
    quelle.write_text("Name;Ort;PLZ\nMüller;Freiburg;79098\n", encoding="utf-8")
    ziel = tmp_path / "deutsch.json"

    tabelle.convert(quelle, ziel, "csv", "json", ctx)

    daten = json.loads(ziel.read_text(encoding="utf-8"))
    assert daten[0]["Ort"] == "Freiburg"


def test_formeln_werden_in_excel_entschaerft(tabelle, tmp_path, ctx):
    """Zellinhalte wie =1+1 duerfen in Excel nicht ausgefuehrt werden."""
    quelle = tmp_path / "boese.csv"
    quelle.write_text('Feld\n"=1+1"\n', encoding="utf-8")
    mappe = tmp_path / "boese.xlsx"

    tabelle.convert(quelle, mappe, "csv", "xlsx", ctx)

    from openpyxl import load_workbook

    blatt = load_workbook(mappe).worksheets[0]
    assert str(blatt.cell(row=2, column=1).value).startswith("'=")


def test_csv_nach_markdown_und_html(tabelle, tmp_path, ctx):
    quelle = tmp_path / "faelle.csv"
    quelle.write_text(BEISPIEL_CSV, encoding="utf-8")

    markdown = tmp_path / "faelle.md"
    tabelle.convert(quelle, markdown, "csv", "md", ctx)
    zeilen = markdown.read_text(encoding="utf-8").splitlines()
    assert zeilen[0].startswith("| Aktenzeichen")
    assert set(zeilen[1].replace(" ", "").replace("|", "")) == {"-"}

    seite = tmp_path / "faelle.html"
    tabelle.convert(quelle, seite, "csv", "html", ctx)
    inhalt = seite.read_text(encoding="utf-8")
    assert "<table>" in inhalt and "Müller GmbH" in inhalt


def test_html_ausgabe_maskiert_sonderzeichen(tabelle, tmp_path, ctx):
    quelle = tmp_path / "boese.csv"
    quelle.write_text('Feld\n"<script>alert(1)</script>"\n', encoding="utf-8")
    ziel = tmp_path / "boese.html"

    tabelle.convert(quelle, ziel, "csv", "html", ctx)

    inhalt = ziel.read_text(encoding="utf-8")
    assert "<script>" not in inhalt
    assert "&lt;script&gt;" in inhalt


def test_csv_nach_sql(tabelle, tmp_path, ctx):
    quelle = tmp_path / "faelle.csv"
    quelle.write_text("Name,Ort\nO'Brien,Freiburg\n", encoding="utf-8")
    ziel = tmp_path / "faelle.sql"

    tabelle.convert(quelle, ziel, "csv", "sql", ctx)

    inhalt = ziel.read_text(encoding="utf-8")
    assert "CREATE TABLE" in inhalt
    # Einfache Anfuehrungszeichen muessen verdoppelt sein.
    assert "'O''Brien'" in inhalt


def test_json_ohne_tabelle_meldet_klaren_fehler(tabelle, tmp_path, ctx):
    quelle = tmp_path / "wert.json"
    quelle.write_text('"nur ein Text"', encoding="utf-8")

    with pytest.raises(ConversionError) as fehler:
        tabelle.convert(quelle, tmp_path / "x.csv", "json", "csv", ctx)
    assert "keine Tabelle" in fehler.value.message


def test_json_nach_yaml_und_zurueck(baum, tmp_path, ctx):
    quelle = tmp_path / "konfig.json"
    quelle.write_text(
        json.dumps({"dienst": {"port": 8080, "namen": ["a", "b"], "aktiv": True}}),
        encoding="utf-8",
    )
    yaml_datei = tmp_path / "konfig.yaml"
    zurueck = tmp_path / "zurueck.json"

    baum.convert(quelle, yaml_datei, "json", "yaml", ctx)
    assert "port: 8080" in yaml_datei.read_text(encoding="utf-8")

    baum.convert(yaml_datei, zurueck, "yaml", "json", ctx)
    assert json.loads(zurueck.read_text(encoding="utf-8"))["dienst"]["port"] == 8080


def test_xml_mit_entitaeten_wird_abgewiesen(baum, tmp_path, ctx):
    """Schutz vor der 'Billion Laughs'-Angriffsform."""
    boesartig = """<?xml version="1.0"?>
<!DOCTYPE lolz [
 <!ENTITY lol "lol">
 <!ENTITY lol1 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">
]>
<root>&lol1;</root>"""
    quelle = tmp_path / "bombe.xml"
    quelle.write_text(boesartig, encoding="utf-8")

    with pytest.raises(ConversionError):
        baum.convert(quelle, tmp_path / "x.json", "xml", "json", ctx)


def test_xml_nach_json(baum, tmp_path, ctx):
    quelle = tmp_path / "akte.xml"
    quelle.write_text(
        '<?xml version="1.0" encoding="UTF-8"?>'
        "<akte><nummer>AZ-1</nummer><ort>Freiburg</ort></akte>",
        encoding="utf-8",
    )
    ziel = tmp_path / "akte.json"

    baum.convert(quelle, ziel, "xml", "json", ctx)

    daten = json.loads(ziel.read_text(encoding="utf-8"))
    assert daten["nummer"] == "AZ-1"


def test_ini_nach_toml(baum, tmp_path, ctx):
    quelle = tmp_path / "dienst.ini"
    quelle.write_text("[server]\nport = 8080\nname = konverter\n", encoding="utf-8")
    ziel = tmp_path / "dienst.toml"

    baum.convert(quelle, ziel, "ini", "toml", ctx)

    inhalt = ziel.read_text(encoding="utf-8")
    assert "[server]" in inhalt and "8080" in inhalt
